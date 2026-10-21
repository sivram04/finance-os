"""The operating model is a real model: forecast cells are formulas, and the statements tie in every scenario."""
import io

import pytest
from openpyxl import load_workbook

from finance_os.app import tools
from finance_os.reports import operating_model


@pytest.fixture(scope="module")
def model():
    data, meta = operating_model.build(tools.Warehouse())
    return data, meta


def test_forecast_cells_are_formulas(model):
    data, meta = model
    wb = load_workbook(io.BytesIO(data))
    for name in ("P&L", "Balance sheet", "Cash flow", "KPIs", "Checks"):
        ws = wb[name]
        forecast_cols = [c for c in range(3, ws.max_column + 1) if ws.cell(6, c).value == "Forecast"]
        assert len(forecast_cols) == operating_model.FORECAST_MONTHS
        for row in range(7, ws.max_row + 1):
            label = ws.cell(row, 1).value or ""
            for c in forecast_cols:
                v = ws.cell(row, c).value
                if v is None or label.startswith("Equity raised"):  # the raise is the one forecast input
                    continue
                assert isinstance(v, str) and v.startswith("="), (name, label, ws.cell(5, c).value, v)


@pytest.mark.parametrize("scenario", ["Base", "Upside", "Downside"])
def test_every_check_passes_in_every_scenario(model, scenario, tmp_path):
    formulas = pytest.importorskip("formulas")
    data, meta = model
    wb = load_workbook(io.BytesIO(data))
    wb["Assumptions"]["C4"] = scenario
    path = tmp_path / "model.xlsx"
    wb.save(path)
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    sheet, cell = meta["status_cell"].split("!")
    status = next(v.value[0][0] for k, v in sol.items() if k.upper().endswith(f"{sheet.upper()}'!{cell}"))
    assert status == "ALL CHECKS PASS"
