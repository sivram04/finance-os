"""Three-statement SaaS operating model (.xlsx), linked with Excel formulas.

    python -m finance_os.reports.operating_model           # -> reports/echo_labs_operating_model.xlsx

Twelve months of actuals from the ledger, then 18 forecast months that are nothing but formulas. The P&L feeds the
balance sheet and the cash flow statement; the cash flow statement's closing cash is the balance sheet's cash; a
scenario switch on Assumptions (Base / Upside / Downside) drives every forecast cell; KPIs (runway, CAC payback)
and integrity checks are formulas too. Nothing in a forecast column is a typed number.

Colour code (banking convention): blue = input (a ledger actual or an assumption), black = formula on the same
sheet, green = formula that links to another sheet.

Sheets: Guide, Assumptions, P&L, Balance sheet, Cash flow, KPIs, Checks, Annual summary, Sensitivity.
"""
import argparse
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from finance_os.app import tools
from finance_os.paths import ROOT

COMPANY = "Echo Labs, Inc."
FORECAST_MONTHS = 18
USD = '#,##0;[Red](#,##0);"–"'
PCT = '0.0%;[Red](0.0%);"–"'
PCT2 = '0.00%;[Red](0.00%);"–"'
NUM1 = '#,##0.0;[Red](#,##0.0);"–"'
INT = '#,##0;[Red](#,##0);"–"'
BLUE, BLACK, GREEN, MUTED = "0000FF", "000000", "008000", "6B7280"
HEAD = PatternFill("solid", fgColor="1F3A5F")
SECTION = PatternFill("solid", fgColor="E8EDF3")
FCST = PatternFill("solid", fgColor="FFF8E6")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")
TOTAL = Border(top=Side(style="thin", color="000000"))
FIRST = 3  # column C: the opening balance month; D.. actuals; then forecast


def _font(color=BLACK, bold=False, italic=False, size=10):
    return Font(name="Calibri", size=size, color=color, bold=bold, italic=italic)


def col(i):
    return get_column_letter(FIRST + i)


def _months_after(m, n):
    y, mo = int(m[:4]), int(m[5:7]) + n
    y += (mo - 1) // 12
    return f"{y}-{(mo - 1) % 12 + 1:02d}"


class Sheet:
    """A statement sheet: rows are named lines, columns are periods (opening, actuals, forecast)."""

    def __init__(self, wb, title, subtitle, periods, n_actual):
        self.ws = wb.create_sheet(title)
        self.title, self.periods, self.n_actual = title, periods, n_actual
        self.rows, self.r = {}, 7
        ws = self.ws
        ws["A1"] = COMPANY
        ws["A1"].font = _font(bold=True, size=13)
        ws["A2"] = title
        ws["A2"].font = _font(bold=True, size=11)
        ws["A3"] = subtitle
        ws["A3"].font = _font(MUTED, italic=True)
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 8
        for i, p in enumerate(periods):
            c = ws.cell(5, FIRST + i, p)
            c.font, c.fill, c.alignment = _font("FFFFFF", bold=True), HEAD, Alignment(horizontal="center")
            kind = "Opening" if i == 0 else "Actual" if i <= n_actual else "Forecast"
            k = ws.cell(6, FIRST + i, kind)
            k.font, k.alignment = _font(MUTED, italic=True), Alignment(horizontal="center")
            ws.column_dimensions[col(i)].width = 13
        ws.cell(5, 1, "USD").font = _font("FFFFFF", bold=True)
        ws.cell(5, 1).fill = HEAD
        ws.cell(5, 2).fill = HEAD
        ws.freeze_panes = "C7"
        ws.sheet_view.showGridLines = False

    def ref(self, name, i, absolute=False, sheet=False):
        c = f"{'$' if absolute else ''}{col(i)}{'$' if absolute else ''}{self.rows[name]}"
        return f"'{self.title}'!{c}" if sheet else c

    def section(self, label):
        c = self.ws.cell(self.r, 1, label)
        c.font, c.fill = _font(bold=True), SECTION
        for i in range(len(self.periods)):
            self.ws.cell(self.r, FIRST + i).fill = SECTION
        self.ws.cell(self.r, 2).fill = SECTION
        self.r += 1

    def line(self, name, label, values, fmt=USD, total=False, indent=1, unit=""):
        """values(i) -> number (an input, blue) or '=formula' (black, green if it reaches another sheet) or None."""
        r = self.r
        self.rows[name] = r
        a = self.ws.cell(r, 1, label)
        a.font, a.alignment = _font(bold=total), Alignment(indent=indent)
        self.ws.cell(r, 2, unit).font = _font(MUTED)
        for i in range(len(self.periods)):
            v = values(i)
            if v is None:
                continue
            c = self.ws.cell(r, FIRST + i, v)
            c.number_format = fmt
            is_formula = isinstance(v, str) and v.startswith("=")
            c.font = _font(GREEN if is_formula and "!" in v else BLACK if is_formula else BLUE, bold=total)
            if i > self.n_actual:
                c.fill = FCST
            if total:
                c.border = TOTAL
        self.r += 1
        return r

    def blank(self):
        self.r += 1


def _drivers(wh, fs):
    """Base-case drivers from the ledger, and the upside / downside around them."""
    p = fs["periods"]
    v = {r["line"] if "line" in r else r["label"]: r for r in fs["rows"]}

    def s(label, months):
        return sum(float(v[label].get(m) or 0) for m in months)

    last3, last6, last12 = p[-3:], p[-6:], p[-12:]
    usage = [float(v["Usage revenue"].get(m) or 0) for m in p]
    growth12 = (usage[-1] / usage[-13]) ** (1 / 12) - 1 if len(usage) > 12 and usage[-13] else 0.0
    growth6 = (usage[-1] / usage[-7]) ** (1 / 6) - 1 if usage[-7] else 0.0
    payroll_lines = ("Salaries and wages", "Payroll taxes", "Benefits")
    payroll_last = sum(float(v[x].get(p[-1]) or 0) for x in payroll_lines)
    nonpay_last = float(v["Total operating expenses"].get(p[-1]) or 0) - payroll_last
    nonpay_first = float(v["Total operating expenses"].get(p[-12]) or 0) - sum(float(v[x].get(p[-12]) or 0) for x in payroll_lines)
    nonpay_growth = (nonpay_last / nonpay_first) ** (1 / 11) - 1 if nonpay_first > 0 else 0.0
    rev_last = float(v["Total revenue"].get(p[-1]) or 0)
    cogs_last = float(v["Total cost of revenue"].get(p[-1]) or 0)
    ar, ap = float(v["Accounts receivable"].get(p[-1]) or 0), float(v["Accounts payable"].get(p[-1]) or 0)
    pay_liab = float(v["Payroll taxes payable"].get(p[-1]) or 0) + float(v["Benefits and deductions payable"].get(p[-1]) or 0)
    sm = wh.query("""select sum(case when department_id = 'SAL' then amount else 0 end) / sum(amount) as share
                     from marts.fct_gl_monthly where account_type = 'opex' and month >= ?::date""", [p[-3] + "-01"])[0]["share"]
    capex12 = -sum(float(v["Capital expenditure"].get(m) or 0) for m in last12)
    gpu = s("COGS - GPU compute", last3) / s("Usage revenue", last3)
    other = (s("Total cost of revenue", last3) - s("COGS - GPU compute", last3)) / s("Total revenue", last3)
    refunds = s("Refunds and credits", last12) / s("Usage revenue", last12)
    dso = ar / rev_last * 30
    dpo = ap / (cogs_last + nonpay_last) * 30
    r4 = lambda x: round(float(x), 4)
    kpi = wh.query("""select max(case when metric_name = 'new_customers' then value end) as new_c,
                             max(case when metric_name = 'logo_churn_rate' then value end) as churn
                      from marts.fct_kpi_monthly where month >= ?::date and month <= ?::date group by month""",
                   [p[-3] + "-01", p[-1] + "-01"])
    new3 = sum(float(r["new_c"] or 0) for r in kpi) / max(len(kpi), 1)
    churn3 = sum(float(r["churn"] or 0) for r in kpi) / max(len(kpi), 1)
    # (key, label, unit, base, upside, downside, source)
    return [
        ("usage_growth", "Usage revenue growth per month", "%", r4(growth12), r4(growth12 + 0.01), r4(growth6 - 0.01),
         "Base: compound growth over the last 12 months; downside: the last 6 months' slowdown persists, minus 1 point"),
        ("services_growth", "Services revenue growth per month", "%", 0.0, 0.01, -0.02, "Project services held flat in the base case"),
        ("refund_pct", "Refunds and credits, % of usage revenue (negative)", "%", r4(refunds), r4(refunds), r4(refunds * 2), "Last 12 months"),
        ("gpu_pct", "GPU compute, % of usage revenue", "%", r4(gpu), r4(gpu - 0.02), r4(gpu + 0.03), "Last 3 months; upside = better GPU utilisation"),
        ("other_cogs_pct", "Other cost of revenue, % of revenue", "%", r4(other), r4(other), r4(other), "Last 3 months: hosting, payment fees, project costs"),
        ("hiring_pct", "Hiring plan executed", "%", 1.0, 1.0, 0.0, "Share of the hiring plan (Assumptions, rows below) that is hired; downside = freeze"),
        ("nonpay_growth", "Non-payroll opex growth per month", "%", r4(max(nonpay_growth, 0.0)), r4(max(nonpay_growth, 0.0) + 0.005),
         r4(max(nonpay_growth, 0.0)), "Base: compound growth over the last 12 months"),
        ("sm_share", "Sales and marketing, % of operating expenses", "%", r4(sm), r4(sm), r4(sm), "Sales department share of opex, last 3 months"),
        ("capex", "Capital expenditure per month", "USD", round(capex12 / 12, 2), round(capex12 / 12, 2), round(capex12 / 24, 2),
         "Average of the last 12 months; downside halves it"),
        ("life", "Useful life of new capex (months)", "months", 36, 36, 36, "Straight line, from the month after purchase"),
        ("dso", "Days sales outstanding", "days", round(dso, 1), round(dso - 5, 1), round(dso + 10, 1), "AR / revenue x 30 at the last close"),
        ("dpo", "Days payable outstanding", "days", round(dpo, 1), round(dpo, 1), round(dpo - 5, 1), "AP / (cost of revenue + non-payroll opex) x 30"),
        ("pay_liab_pct", "Payroll liabilities, % of monthly payroll", "%", r4(pay_liab / payroll_last), r4(pay_liab / payroll_last),
         r4(pay_liab / payroll_last), "Payroll taxes and benefits payable at the last close"),
        ("tax_rate", "Cash tax rate on positive pre-tax income", "%", 0.0, 0.0, 0.0,
         "US losses carried forward shelter profits in the forecast window (Controller page, ASC 740)"),
        ("min_cash", "Minimum cash (treasury policy)", "USD", 3000000, 3000000, 3000000, "seeds.treasury_policy MIN-CASH"),
        ("new_customers", "New customers per month", "count", round(new3, 2), round(new3 * 1.5, 2), round(new3 * 0.5, 2),
         "Last 3 months' average; drives CAC and active customers"),
        ("logo_churn", "Logo churn per month", "%", r4(churn3), r4(churn3 * 0.75), r4(churn3 * 1.25), "Last 3 months' average"),
    ]


def build(wh=None):
    wh = wh or tools.Warehouse()
    fs = tools.get_financial_statements(wh, period="month", periods=13)
    v = {r["line"] if "line" in r else r["label"]: r for r in fs["rows"]}
    actual = fs["periods"]                       # 13 months: the first is the opening balance sheet
    n_act = len(actual) - 1
    fc = [_months_after(actual[-1], k + 1) for k in range(FORECAST_MONTHS)]
    periods = actual + fc
    n = len(periods)
    A = lambda i: 1 <= i <= n_act           # an actual month (has a P&L and cash flow)
    F = lambda i: i > n_act                 # a forecast month
    act = lambda label: (lambda i: float(v[label].get(periods[i]) or 0) if i <= n_act else None)

    wb = Workbook()
    wb.remove(wb.active)
    guide = wb.create_sheet("Guide")

    # ---------------- Assumptions ----------------
    asm = wb.create_sheet("Assumptions")
    asm.sheet_view.showGridLines = False
    asm["A1"], asm["A2"] = COMPANY, "Assumptions"
    asm["A1"].font, asm["A2"].font = _font(bold=True, size=13), _font(bold=True, size=11)
    asm["A3"] = "Change the scenario in C4: every forecast cell on every sheet recalculates."
    asm["A3"].font = _font(MUTED, italic=True)
    asm["A4"], asm["C4"] = "Scenario", "Base"
    asm["A4"].font = _font(bold=True)
    asm["C4"].font, asm["C4"].fill = _font(BLUE, bold=True), INPUT_FILL
    dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False)
    asm.add_data_validation(dv)
    dv.add("C4")
    hdr = ["Driver", "Unit", "Base", "Upside", "Downside", "Live (selected)", "Source"]
    for j, h in enumerate(hdr, 1):
        c = asm.cell(6, j, h)
        c.font, c.fill = _font("FFFFFF", bold=True), HEAD
    drivers = _drivers(wh, fs)  # all 13 months, so growth can be measured over 12
    D = {}
    for k, (key, label, unit, base, up, down, src) in enumerate(drivers):
        r = 7 + k
        asm.cell(r, 1, label)
        asm.cell(r, 2, unit).font = _font(MUTED)
        fmt = PCT2 if unit == "%" else USD if unit == "USD" else NUM1
        for j, val in ((3, base), (4, up), (5, down)):
            c = asm.cell(r, j, val)
            c.font, c.number_format = _font(BLUE), fmt
        pick = f"INDEX(C{r}:E{r},MATCH($C$4,$C$6:$E$6,0))"
        override = {"usage_growth": "Sensitivity!$C$5", "nonpay_growth": "Sensitivity!$C$6"}.get(key)
        live = asm.cell(r, 6, f'=IF({override}="",{pick},{override})' if override else f"={pick}")
        live.font, live.number_format = _font(bold=True), fmt
        asm.cell(r, 7, src).font = _font(MUTED, italic=True)
        D[key] = f"Assumptions!$F${r}"
    for j, w in zip("ABCDEFG", (44, 8, 12, 12, 12, 15, 70)):
        asm.column_dimensions[j].width = w

    # hiring plan by forecast month (inputs), from seeds.hiring_plan
    plan = wh.query("""select strftime(planned_start, '%Y-%m') as m, count(*) as hires,
                              sum(annual_salary * (1 + p.employer_tax_pct) / 12 + p.benefits_per_month) as monthly_cost
                       from seeds.hiring_plan h cross join seeds.labor_cost_policy p
                       where h.status in ('open', 'approved') group by 1""")
    plan = {r["m"]: r for r in plan}
    hr = 8 + len(drivers)
    asm.cell(hr, 1, "Hiring plan (open and approved roles)").font = _font(bold=True)
    asm.cell(hr + 1, 1, "Month")
    asm.cell(hr + 2, 1, "New hires")
    asm.cell(hr + 3, 1, "Monthly loaded cost added (salary + 7.65% payroll tax + $950 benefits)")
    HIRE = {}
    for k, m in enumerate(fc):
        cc = get_column_letter(3 + k)
        asm.cell(hr + 1, 3 + k, m).font = _font(bold=True)
        h = asm.cell(hr + 2, 3 + k, int(plan.get(m, {}).get("hires") or 0))
        cst = asm.cell(hr + 3, 3 + k, round(float(plan.get(m, {}).get("monthly_cost") or 0), 2))
        h.font = cst.font = _font(BLUE)
        cst.number_format = USD
        HIRE[m] = (f"Assumptions!${cc}${hr + 2}", f"Assumptions!${cc}${hr + 3}")
        asm.column_dimensions[cc].width = max(asm.column_dimensions[cc].width or 0, 12)
    asm.freeze_panes = "A7"

    # ---------------- P&L ----------------
    pl = Sheet(wb, "P&L", "Monthly income statement · actuals from the ledger (blue), forecast by formula (shaded) · USD", periods, n_act)
    PL, BS, CF = "P&L", "Balance sheet", "Cash flow"
    prev = lambda i: col(i - 1)
    pl.section("Revenue")
    pl.line("usage", "Usage revenue", lambda i: act("Usage revenue")(i) if A(i) else
            f"={prev(i)}{pl.r}*(1+{D['usage_growth']})" if F(i) else None)
    pl.line("refunds", "Refunds and credits", lambda i: act("Refunds and credits")(i) if A(i) else
            f"={col(i)}{pl.rows['usage']}*{D['refund_pct']}" if F(i) else None)
    pl.line("services", "Services revenue", lambda i: act("Services revenue")(i) if A(i) else
            f"={prev(i)}{pl.r}*(1+{D['services_growth']})" if F(i) else None)
    pl.line("revenue", "Total revenue", lambda i: f"=SUM({col(i)}{pl.rows['usage']}:{col(i)}{pl.rows['services']})" if i else None, total=True)
    pl.section("Cost of revenue")
    pl.line("gpu", "GPU compute", lambda i: act("COGS - GPU compute")(i) if A(i) else
            f"={col(i)}{pl.rows['usage']}*{D['gpu_pct']}" if F(i) else None)
    pl.line("other_cogs", "Hosting, payment fees and project costs",
            lambda i: (act("Total cost of revenue")(i) - act("COGS - GPU compute")(i)) if A(i) else
            f"={col(i)}{pl.rows['revenue']}*{D['other_cogs_pct']}" if F(i) else None)
    pl.line("cogs", "Total cost of revenue", lambda i: f"={col(i)}{pl.rows['gpu']}+{col(i)}{pl.rows['other_cogs']}" if i else None, total=True)
    pl.line("gross_profit", "Gross profit", lambda i: f"={col(i)}{pl.rows['revenue']}-{col(i)}{pl.rows['cogs']}" if i else None, total=True)
    pl.line("gm", "Gross margin", lambda i: f"=IF({col(i)}{pl.rows['revenue']}=0,0,{col(i)}{pl.rows['gross_profit']}/{col(i)}{pl.rows['revenue']})" if i else None,
            fmt=PCT, indent=2)
    pl.section("Operating expenses")
    pay_act = lambda i: sum(act(x)(i) for x in ("Salaries and wages", "Payroll taxes", "Benefits"))
    pl.line("payroll", "Payroll (salaries, payroll taxes, benefits)", lambda i: pay_act(i) if A(i) else
            f"={prev(i)}{pl.r}+{HIRE[periods[i]][1]}*{D['hiring_pct']}" if F(i) else None)
    pl.line("nonpayroll", "Non-payroll opex (cloud R&D, software, rent, fees, marketing)",
            lambda i: (act("Total operating expenses")(i) - pay_act(i)) if A(i) else
            f"={prev(i)}{pl.r}*(1+{D['nonpay_growth']})" if F(i) else None)
    pl.line("opex", "Total operating expenses", lambda i: f"={col(i)}{pl.rows['payroll']}+{col(i)}{pl.rows['nonpayroll']}" if i else None, total=True)
    pl.line("ebitda", "EBITDA", lambda i: f"={col(i)}{pl.rows['gross_profit']}-{col(i)}{pl.rows['opex']}" if i else None, total=True)
    pl.line("dep", "Depreciation", lambda i: act("Depreciation")(i) if A(i) else
            f"=${col(n_act)}${pl.r}+SUM(${col(n_act + 1)}${{cap}}:{prev(i)}${{cap}})/{D['life']}" if F(i) else None)
    pl.line("ebit", "Operating income (EBIT)", lambda i: f"={col(i)}{pl.rows['ebitda']}-{col(i)}{pl.rows['dep']}" if i else None, total=True)
    pl.line("tax", "Income tax", lambda i: 0 if A(i) else f"=MAX(0,{col(i)}{pl.rows['ebit']})*{D['tax_rate']}" if F(i) else None)
    pl.line("net_income", "Net income", lambda i: f"={col(i)}{pl.rows['ebit']}-{col(i)}{pl.rows['tax']}" if i else None, total=True)
    pl.blank()
    pl.section("Memo")
    pl.line("ledger_ni", "Net income per the ledger (actuals)", lambda i: act("Net income")(i) if A(i) else None)
    hc = _headcount(wh, periods[: n_act + 1])
    pl.line("headcount", "Headcount (month end)", lambda i: hc[i] if i <= n_act else
            f"={prev(i)}{pl.r}+{HIRE[periods[i]][0]}*{D['hiring_pct']}", fmt=NUM1)

    # ---------------- Cash flow (rows reserved first so the balance sheet can link to it) ----------------
    bs = Sheet(wb, BS, "Month-end balance sheet · actuals from the ledger (blue), forecast by formula (shaded) · USD", periods, n_act)
    cf = Sheet(wb, CF, "Indirect method, linked to the P&L and the balance sheet for every month (actuals included) · USD", periods, n_act)
    cap_row = None

    bs.section("Assets")
    bs.line("cash", "Cash", lambda i: act("Cash - operating")(i) if i <= n_act else f"='{CF}'!{col(i)}{{cf_close}}")
    bs.line("ar", "Accounts receivable", lambda i: act("Accounts receivable")(i) if i <= n_act else
            f"='{PL}'!{col(i)}{pl.rows['revenue']}*{D['dso']}/30")
    bs.line("fa_cost", "Fixed assets at cost", lambda i: act("Fixed assets at cost")(i) if i <= n_act else
            f"={prev(i)}{bs.r}-'{CF}'!{col(i)}{{cf_capex}}")
    bs.line("acc_dep", "Accumulated depreciation", lambda i: act("Accumulated depreciation")(i) if i <= n_act else
            f"={prev(i)}{bs.r}-'{PL}'!{col(i)}{pl.rows['dep']}")
    bs.line("assets", "Total assets", lambda i: f"=SUM({col(i)}{bs.rows['cash']}:{col(i)}{bs.rows['acc_dep']})", total=True)
    bs.section("Liabilities")
    bs.line("ap", "Accounts payable", lambda i: act("Accounts payable")(i) if i <= n_act else
            f"=('{PL}'!{col(i)}{pl.rows['cogs']}+'{PL}'!{col(i)}{pl.rows['nonpayroll']})*{D['dpo']}/30")
    bs.line("pay_liab", "Payroll taxes and benefits payable",
            lambda i: act("Payroll taxes payable")(i) + act("Benefits and deductions payable")(i) if i <= n_act else
            f"='{PL}'!{col(i)}{pl.rows['payroll']}*{D['pay_liab_pct']}")
    bs.line("liabilities", "Total liabilities", lambda i: f"={col(i)}{bs.rows['ap']}+{col(i)}{bs.rows['pay_liab']}", total=True)
    bs.section("Equity")
    bs.line("paid_in", "Preferred stock and paid-in capital", lambda i: act("Equity - preferred stock")(i) if i <= n_act else
            f"={prev(i)}{bs.r}+'{CF}'!{col(i)}{{cf_equity}}")
    bs.line("re", "Retained earnings", lambda i: act("Retained earnings")(i) if i <= n_act else
            f"={prev(i)}{bs.r}+'{PL}'!{col(i)}{pl.rows['net_income']}")
    bs.line("equity", "Total equity", lambda i: f"={col(i)}{bs.rows['paid_in']}+{col(i)}{bs.rows['re']}", total=True)
    bs.line("le", "Total liabilities and equity", lambda i: f"={col(i)}{bs.rows['liabilities']}+{col(i)}{bs.rows['equity']}", total=True)
    bs.blank()
    bs.line("check", "Check: assets - liabilities - equity (must be 0)",
            lambda i: f"=ROUND({col(i)}{bs.rows['assets']}-{col(i)}{bs.rows['le']},2)", indent=1)

    B = lambda name, i: f"'{BS}'!{col(i)}{bs.rows[name]}"
    cf.section("Operating activities")
    cf.line("ni", "Net income", lambda i: f"='{PL}'!{col(i)}{pl.rows['net_income']}" if i else None)
    cf.line("dep", "Depreciation (non-cash, added back)", lambda i: f"='{PL}'!{col(i)}{pl.rows['dep']}" if i else None)
    cf.line("d_ar", "Change in accounts receivable", lambda i: f"={B('ar', i - 1)}-{B('ar', i)}" if i else None)
    cf.line("d_ap", "Change in accounts payable", lambda i: f"={B('ap', i)}-{B('ap', i - 1)}" if i else None)
    cf.line("d_pay", "Change in payroll liabilities", lambda i: f"={B('pay_liab', i)}-{B('pay_liab', i - 1)}" if i else None)
    cf.line("cfo", "Cash from operations", lambda i: f"=SUM({col(i)}{cf.rows['ni']}:{col(i)}{cf.rows['d_pay']})" if i else None, total=True)
    cf.section("Investing activities")
    cf.line("capex", "Capital expenditure", lambda i: f"=-({B('fa_cost', i)}-{B('fa_cost', i - 1)})" if A(i) else
            f"=-{D['capex']}" if F(i) else None)
    cf.section("Financing activities")
    cf.line("equity", "Equity raised (enter a raise here)", lambda i: f"={B('paid_in', i)}-{B('paid_in', i - 1)}" if A(i) else 0 if F(i) else None)
    cf.blank()
    cf.line("net", "Net change in cash", lambda i: f"={col(i)}{cf.rows['cfo']}+{col(i)}{cf.rows['capex']}+{col(i)}{cf.rows['equity']}" if i else None, total=True)
    cf.line("open", "Opening cash", lambda i: f"={B('cash', i - 1)}" if i else None)
    cf.line("close", "Closing cash", lambda i: f"={col(i)}{cf.rows['open']}+{col(i)}{cf.rows['net']}" if i else None, total=True)
    cf.line("fcf", "Free cash flow (operations + capex)", lambda i: f"={col(i)}{cf.rows['cfo']}+{col(i)}{cf.rows['capex']}" if i else None)

    # fill the forward references (cash flow rows used by the balance sheet and P&L depreciation)
    fills = {"{cf_close}": str(cf.rows["close"]), "{cf_capex}": str(cf.rows["capex"]), "{cf_equity}": str(cf.rows["equity"])}
    for ws in (bs.ws,):
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    for k2, val in fills.items():
                        c.value = c.value.replace(k2, val)
    # depreciation on new capex: cumulative forecast capex (a negative cash flow) over the useful life
    dep_r = pl.rows["dep"]
    for i in range(n_act + 1, n):
        c = pl.ws.cell(dep_r, FIRST + i)
        existing = f"${col(n_act)}${dep_r}"
        c.value = (f"={existing}" if i == n_act + 1 else
                   f"={existing}-SUM('{CF}'!${col(n_act + 1)}${cf.rows['capex']}:'{CF}'!{col(i - 1)}{cf.rows['capex']})/{D['life']}")
        c.font = _font(GREEN)

    # ---------------- KPIs ----------------
    kp = Sheet(wb, "KPIs", "Runway, burn and unit economics, all by formula · USD unless stated", periods, n_act)
    P = lambda name, i: f"'{PL}'!{col(i)}{pl.rows[name]}"
    kp.section("Cash")
    kp.line("cash", "Cash at month end", lambda i: f"={B('cash', i)}")
    kp.line("burn", "Net burn (cash used in the month)", lambda i: f"=-'{CF}'!{col(i)}{cf.rows['net']}" if i else None)
    kp.line("burn3", "Average net burn, last 3 months", lambda i: f"=AVERAGE({col(i - 2)}{kp.r - 1}:{col(i)}{kp.r - 1})" if i >= 3 else None)
    kp.line("runway", "Runway (months) at the 3-month burn rate",
            lambda i: f'=IF({col(i)}{kp.rows["burn3"]}<=0,"cash generative",{col(i)}{kp.rows["cash"]}/{col(i)}{kp.rows["burn3"]})' if i >= 3 else None, fmt=NUM1)
    kp.line("headroom", "Cash above the minimum cash policy", lambda i: f"={col(i)}{kp.rows['cash']}-{D['min_cash']}")
    kp.section("Customers and unit economics (same definitions as Finance OS)")
    sm_act = {str(r["m"]): r for r in wh.query("""select strftime(s.month, '%Y-%m') as m, s.cac_payback_months,
                                                         k.active_customers, k.new_customers, k.logo_churn_rate,
                                                         (select sum(natural_amount) from marts.fct_gl_monthly g
                                                          where g.month = s.month and g.account_type = 'opex' and g.department_id = 'SAL') as sm_spend
                                                  from marts.fct_saas_metrics_monthly s
                                                  join (select month,
                                                               max(case when metric_name = 'active_customers' then value end) as active_customers,
                                                               max(case when metric_name = 'new_customers' then value end) as new_customers,
                                                               max(case when metric_name = 'logo_churn_rate' then value end) as logo_churn_rate
                                                        from marts.fct_kpi_monthly group by 1) k on k.month = s.month""")}
    ua = lambda key: (lambda i: float(sm_act.get(periods[i], {}).get(key) or 0))
    kp.line("arr", "ARR (usage revenue x 12)", lambda i: f"={P('usage', i)}*12" if i else None)
    kp.line("sm", "Sales and marketing spend (Sales department)", lambda i: ua("sm_spend")(i) if A(i) else
            f"={P('opex', i)}*{D['sm_share']}" if F(i) else None)
    kp.line("new_c", "New customers", lambda i: ua("new_customers")(i) if A(i) else f"={D['new_customers']}" if F(i) else None, fmt=NUM1)
    kp.line("churn", "Logo churn rate", lambda i: ua("logo_churn_rate")(i) if A(i) else f"={D['logo_churn']}" if F(i) else None, fmt=PCT)
    kp.line("active", "Active customers (month end)", lambda i: ua("active_customers")(i) if A(i) else
            f"={prev(i)}{kp.r}*(1-{col(i)}{kp.rows['churn']})+{col(i)}{kp.rows['new_c']}" if F(i) else None, fmt=NUM1)
    kp.line("arpa", "ARPA (revenue per active customer, monthly)", lambda i: f"=IF({col(i)}{kp.rows['active']}=0,0,{P('revenue', i)}/{col(i)}{kp.rows['active']})" if i else None)
    t3 = lambda name, i: f"SUM({col(i - 2)}{kp.rows[name]}:{col(i)}{kp.rows[name]})"
    kp.line("cac", "CAC (S&M spend / new customers, trailing 3 months)",
            lambda i: f"=IF({t3('new_c', i)}=0,0,{t3('sm', i)}/{t3('new_c', i)})" if i >= 3 else None)
    kp.line("cac_payback", "CAC payback (months) = CAC / (ARPA x gross margin)",
            lambda i: f"=IF({col(i)}{kp.rows['arpa']}*{P('gm', i)}=0,0,{col(i)}{kp.rows['cac']}/({col(i)}{kp.rows['arpa']}*{P('gm', i)}))" if i >= 3 else None, fmt=NUM1)
    kp.line("ltv", "LTV (ARPA x gross margin / churn, trailing 3 months)",
            lambda i: f"=IF(AVERAGE({col(i - 2)}{kp.rows['churn']}:{col(i)}{kp.rows['churn']})=0,0,{col(i)}{kp.rows['arpa']}*{P('gm', i)}/AVERAGE({col(i - 2)}{kp.rows['churn']}:{col(i)}{kp.rows['churn']}))" if i >= 3 else None)
    kp.line("ltv_cac", "LTV : CAC", lambda i: f"=IF({col(i)}{kp.rows['cac']}=0,0,{col(i)}{kp.rows['ltv']}/{col(i)}{kp.rows['cac']})" if i >= 3 else None, fmt='0.00"x"')
    kp.line("app_payback", "CAC payback per Finance OS (actual months, for the tie-out)",
            lambda i: ua("cac_payback_months")(i) if A(i) and i >= 3 else None, fmt=NUM1)
    kp.line("gm", "Gross margin", lambda i: f"={P('gm', i)}" if i else None, fmt=PCT)
    kp.line("ebitda_margin", "EBITDA margin", lambda i: f"=IF({P('revenue', i)}=0,0,{P('ebitda', i)}/{P('revenue', i)})" if i else None, fmt=PCT)

    # ---------------- Checks ----------------
    ck = Sheet(wb, "Checks", "Every check is a formula; TRUE in every month means the model is internally consistent", periods, n_act)
    ck.section("Integrity checks (TRUE = pass)")
    ck.line("bs", "Balance sheet balances (assets = liabilities + equity)", lambda i: f"=ABS({B('check', i)})<0.01", fmt="General")
    ck.line("cash", "Cash flow closing cash = balance sheet cash", lambda i: f"=ABS('{CF}'!{col(i)}{cf.rows['close']}-{B('cash', i)})<0.01" if i else None, fmt="General")
    ck.line("ni", "Cash flow starts from P&L net income", lambda i: f"=ABS('{CF}'!{col(i)}{cf.rows['ni']}-{P('net_income', i)})<0.01" if i else None, fmt="General")
    ck.line("ledger", "P&L net income ties to the ledger (actual months)", lambda i: f"=ABS({P('net_income', i)}-{P('ledger_ni', i)})<0.01" if A(i) else None, fmt="General")
    ck.line("cac", "CAC payback ties to Finance OS (actual months)",
            lambda i: f"=ABS(ROUND(KPIs!{col(i)}{kp.rows['cac_payback']},1)-KPIs!{col(i)}{kp.rows['app_payback']})<0.15" if A(i) and i >= 3 else None, fmt="General")
    ck.line("re", "Retained earnings roll forward by net income", lambda i: f"=ABS({B('re', i)}-{B('re', i - 1)}-{P('net_income', i)})<0.01" if i else None, fmt="General")
    ck.line("month_ok", "All checks pass this month", lambda i: f"=AND({col(i)}{ck.rows['bs']}:{col(i)}{ck.r - 1})" if i else f"={col(i)}{ck.rows['bs']}", fmt="General", total=True)
    ck.blank()
    r = ck.r
    ck.ws.cell(r, 1, "Model status").font = _font(bold=True)
    last = col(n - 1)
    st = ck.ws.cell(r, FIRST, f'=IF(COUNTIF({col(0)}{ck.rows["month_ok"]}:{last}{ck.rows["month_ok"]},FALSE)=0,"ALL CHECKS PASS","CHECK FAILED")')
    st.font = _font(bold=True)
    ck.ws.cell(r + 1, 1, "Months failing a check")
    ck.ws.cell(r + 1, FIRST, f'=COUNTIF({col(0)}{ck.rows["month_ok"]}:{last}{ck.rows["month_ok"]},FALSE)')
    ck.rows["status"], ck.rows["failing"] = r, r + 1
    for name in ("bs", "cash", "ni", "ledger", "cac", "re", "month_ok"):
        rng = f"{col(0)}{ck.rows[name]}:{last}{ck.rows[name]}"
        ck.ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{col(0)}{ck.rows[name]}=FALSE"], fill=PatternFill("solid", fgColor="F8D7DA")))
        ck.ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{col(0)}{ck.rows[name]}=TRUE"], fill=PatternFill("solid", fgColor="D4EDDA")))

    # ---------------- Annual summary ----------------
    an = wb.create_sheet("Annual summary")
    an.sheet_view.showGridLines = False
    an["A1"], an["A2"] = COMPANY, "Annual summary"
    an["A1"].font, an["A2"].font = _font(bold=True, size=13), _font(bold=True, size=11)
    an["A3"] = "Full calendar years in the model; each figure sums or picks the monthly columns, so it follows the scenario."
    an["A3"].font = _font(MUTED, italic=True)
    years = sorted({periods[i][:4] for i in range(1, n) if sum(1 for j in range(1, n) if periods[j][:4] == periods[i][:4]) == 12})
    an.column_dimensions["A"].width = 44
    for k, y in enumerate(years):
        ycols = [i for i in range(1, n) if periods[i][:4] == y]
        nact = sum(1 for i in ycols if A(i))
        c = an.cell(5, 3 + k, f"FY{y}")
        c.font, c.fill, c.alignment = _font("FFFFFF", bold=True), HEAD, Alignment(horizontal="center")
        an.cell(6, 3 + k, f"{nact} actual + {12 - nact} forecast" if nact else "forecast").font = _font(MUTED, italic=True)
        an.column_dimensions[get_column_letter(3 + k)].width = 18
    an.cell(5, 1, "USD").font = _font("FFFFFF", bold=True)
    an.cell(5, 1).fill = HEAD
    an.cell(5, 2).fill = HEAD

    def ysum(sheet, row, y):
        ycols = [i for i in range(1, n) if periods[i][:4] == y]
        return f"=SUM('{sheet}'!{col(ycols[0])}{row}:{col(ycols[-1])}{row})"

    def yend(sheet, row, y):
        return f"='{sheet}'!{col(max(i for i in range(1, n) if periods[i][:4] == y))}{row}"

    lines = [("Revenue", lambda y: ysum(PL, pl.rows["revenue"], y), USD, True),
             ("Gross profit", lambda y: ysum(PL, pl.rows["gross_profit"], y), USD, False),
             ("Gross margin", None, PCT, False),
             ("Operating expenses", lambda y: ysum(PL, pl.rows["opex"], y), USD, False),
             ("EBITDA", lambda y: ysum(PL, pl.rows["ebitda"], y), USD, True),
             ("EBITDA margin", None, PCT, False),
             ("Net income", lambda y: ysum(PL, pl.rows["net_income"], y), USD, True),
             ("Cash from operations", lambda y: ysum(CF, cf.rows["cfo"], y), USD, False),
             ("Capital expenditure", lambda y: ysum(CF, cf.rows["capex"], y), USD, False),
             ("Free cash flow", lambda y: ysum(CF, cf.rows["fcf"], y), USD, True),
             ("Cash at year end", lambda y: yend(BS, bs.rows["cash"], y), USD, True),
             ("ARR at year end", lambda y: yend("KPIs", kp.rows["arr"], y), USD, False),
             ("Active customers at year end", lambda y: yend("KPIs", kp.rows["active"], y), NUM1, False),
             ("Headcount at year end", lambda y: yend(PL, pl.rows["headcount"], y), NUM1, False),
             ("CAC payback at year end (months)", lambda y: yend("KPIs", kp.rows["cac_payback"], y), NUM1, False)]
    arow = {}
    for k, (label, fn, fmt, bold) in enumerate(lines):
        r = 7 + k
        arow[label] = r
        an.cell(r, 1, label).font = _font(bold=bold)
        for j, y in enumerate(years):
            cl = get_column_letter(3 + j)
            if label == "Gross margin":
                f = f"=IF({cl}{arow['Revenue']}=0,0,{cl}{arow['Gross profit']}/{cl}{arow['Revenue']})"
            elif label == "EBITDA margin":
                f = f"=IF({cl}{arow['Revenue']}=0,0,{cl}{arow['EBITDA']}/{cl}{arow['Revenue']})"
            else:
                f = fn(y)
            c = an.cell(r, 3 + j, f)
            c.number_format, c.font = fmt, _font(GREEN if "!" in f else BLACK, bold=bold)
    an.freeze_panes = "C7"

    # ---------------- Sensitivity (an Excel Data Table) ----------------
    from openpyxl.worksheet.formula import DataTableFormula
    se = wb.create_sheet("Sensitivity")
    se.sheet_view.showGridLines = False
    se["A1"], se["A2"] = COMPANY, "Sensitivity: cash at the end of the forecast"
    se["A1"].font, se["A2"].font = _font(bold=True, size=13), _font(bold=True, size=11)
    se["A3"] = ("An Excel Data Table (Data > What-If Analysis > Data Table): Excel reruns the whole model for each pair of "
                "usage growth (down) and non-payroll opex growth (across), with the selected scenario's other drivers. "
                "Leave C5 and C6 blank: they are the table's input cells.")
    se["A3"].font, se["A3"].alignment = _font(MUTED, italic=True), Alignment(wrap_text=True)
    se.merge_cells("A3:H3")
    se.row_dimensions[3].height = 44
    se["A5"] = "Usage growth override (the table's column input cell)"
    se["A6"] = "Non-payroll opex growth override (the table's row input cell)"
    for c in ("C5", "C6"):
        se[c].fill, se[c].font, se[c].number_format = INPUT_FILL, _font(BLUE), PCT2
    end_cash = f"='{BS}'!{col(n - 1)}{bs.rows['cash']}"
    se["A8"] = "Cash at the end of the forecast (down: usage growth per month; across: non-payroll opex growth per month)"
    se["A8"].font = _font(bold=True)
    se["B9"] = end_cash
    se["B9"].font, se["B9"].number_format = _font(GREEN), USD
    g_vals = [0.0, 0.01, 0.02, 0.03, 0.04]
    o_vals = [0.0, 0.01, 0.02, 0.03]
    for j, o in enumerate(o_vals):
        c = se.cell(9, 3 + j, o)
        c.font, c.number_format, c.fill = _font(BLUE, bold=True), PCT, SECTION
    for k, g in enumerate(g_vals):
        c = se.cell(10 + k, 2, g)
        c.font, c.number_format, c.fill = _font(BLUE, bold=True), PCT, SECTION
    table_ref = f"C10:{get_column_letter(2 + len(o_vals))}{9 + len(g_vals)}"
    se["C10"] = DataTableFormula(ref=table_ref, dt2D=True, r1="C6", r2="C5")
    for k in range(len(g_vals)):
        for j in range(len(o_vals)):
            se.cell(10 + k, 3 + j).number_format = USD
    se.column_dimensions["A"].width = 64
    se.column_dimensions["B"].width = 16
    for j in range(len(o_vals)):
        se.column_dimensions[get_column_letter(3 + j)].width = 15
    se["A16"] = ("Google Sheets does not support data tables, so there the grid stays empty; the Finance OS web view fills "
                 "it by recalculating the model for each pair.")
    se["A16"].font = _font(MUTED, italic=True)
    sens = {"table": table_ref, "growth": g_vals, "opex": o_vals, "row_input": "C6", "col_input": "C5", "output": end_cash[1:]}

    # ---------------- summary on Assumptions (live scenario outputs) ----------------
    sr = hr + 5
    asm.cell(sr, 1, "Outputs for the selected scenario").font = _font(bold=True)
    end = col(n - 1)
    outs = [("Revenue, last forecast month", f"='{PL}'!{end}{pl.rows['revenue']}", USD),
            ("Lowest month-end cash in the forecast", f"=MIN('{BS}'!{col(n_act + 1)}{bs.rows['cash']}:{end}{bs.rows['cash']})", USD),
            ("Cash at the end of the forecast", f"='{BS}'!{end}{bs.rows['cash']}", USD),
            ("Runway at the end of the forecast (months)", f"=KPIs!{end}{kp.rows['runway']}", NUM1),
            ("CAC payback at the end of the forecast (months)", f"=KPIs!{end}{kp.rows['cac_payback']}", NUM1),
            ("EBITDA margin, last forecast month", f"=KPIs!{end}{kp.rows['ebitda_margin']}", PCT),
            ("CAC payback, last actual month (months)", f"=KPIs!{col(n_act)}{kp.rows['cac_payback']}", NUM1),
            ("Model checks", f"=Checks!{col(0)}{ck.rows['status']}", "General")]
    for k, (label, f, fmt) in enumerate(outs):
        asm.cell(sr + 1 + k, 1, label)
        c = asm.cell(sr + 1 + k, 3, f)
        c.font, c.number_format = _font(GREEN, bold=True), fmt

    _guide(guide, actual, fc)
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates everything, including the data table, on open
    wb.active = wb.sheetnames.index("Assumptions")
    buf = io.BytesIO()
    wb.save(buf)
    meta = {"actual_from": actual[1], "actual_to": actual[-1], "forecast_from": fc[0], "forecast_to": fc[-1],
            "status_cell": f"Checks!{col(0)}{ck.rows['status']}", "outputs_row": sr + 1, "outputs": len(outs),
            "scenario_cell": "Assumptions!C4", "sensitivity": sens}
    return buf.getvalue(), meta


_GRID_CACHE = {}


def engine_copy(wb, path, scenario="Base"):
    """Save a copy of the workbook the formula engine can load: the scenario set, the Excel data table removed
    (Excel computes it on open; the engine can't read data tables, so sensitivity() reruns the model instead)."""
    wb["Assumptions"]["C4"] = scenario
    if "Sensitivity" in wb.sheetnames:
        wb["Sensitivity"]["C10"] = None
    wb.save(path)


def sensitivity(xl, meta):
    """The data table's values, by recalculating the model for each pair of overrides (what Excel's data table does)."""
    sens = meta["sensitivity"]
    key = lambda cell: next(k for k in xl.dsp.data_nodes if k.upper().endswith(f"SENSITIVITY'!{cell}"))
    sheet, cell = sens["output"].split("!")
    out = next(k for k in xl.dsp.data_nodes if k.upper().endswith(f"{sheet.strip(chr(39)).upper()}'!{cell}"))
    grid = []
    for g in sens["growth"]:
        row = []
        for o in sens["opex"]:
            sol = xl.calculate(inputs={key(sens["col_input"]): g, key(sens["row_input"]): o}, outputs=[out])
            v = sol[out].value[0][0]
            row.append(float(v.item() if hasattr(v, "item") else v))
        grid.append(row)
    return grid


def grid(wh, scenario="Base"):
    """The model as the app shows it: for every sheet and cell, the formula (or the input), the value Excel would
    calculate for the chosen scenario, its number format and its role (input / formula / link). Calculated with a
    formula engine from the same workbook the Excel download contains."""
    import tempfile
    import formulas
    from openpyxl import load_workbook
    scenario = scenario if scenario in ("Base", "Upside", "Downside") else "Base"
    stamp = str(wh.query("select max(entry_date) as d from staging.stg_erp__journal_entries")[0]["d"])
    key = (scenario, stamp)
    if key in _GRID_CACHE:
        return _GRID_CACHE[key]
    data, meta = build(wh)
    wb = load_workbook(io.BytesIO(data))
    with tempfile.TemporaryDirectory() as d:
        path = f"{d}/model.xlsx"
        engine_copy(wb, path, scenario)
        xl = formulas.ExcelModel().loads(path).finish()
        sol = xl.calculate()
        sens_grid = sensitivity(xl, meta)
    wb = load_workbook(io.BytesIO(data))
    wb["Assumptions"]["C4"] = scenario
    values = {}
    for k, v in sol.items():
        if "!" not in k or ":" in k.split("!")[1]:
            continue
        sheet, cell = k.split("!")
        sheet = sheet.split("]")[-1].strip("'").upper()
        try:
            val = v.value[0][0]
        except Exception:  # noqa: BLE001
            continue
        val = val.item() if hasattr(val, "item") else val
        values[(sheet, cell.replace("$", ""))] = None if str(val) in ("empty", "") else val
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    first_col, first_row = coordinate_from_string(meta["sensitivity"]["table"].split(":")[0])
    table_f = f'{{=TABLE({meta["sensitivity"]["row_input"]},{meta["sensitivity"]["col_input"]})}}'
    for k, row in enumerate(sens_grid):
        for j, val in enumerate(row):
            ref = f"{get_column_letter(column_index_from_string(first_col) + j)}{first_row + k}"
            values[("SENSITIVITY", ref)] = val
            wb["Sensitivity"][ref] = "=TABLE"
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None:
                    continue
                f = c.value if isinstance(c.value, str) and c.value.startswith("=") else None
                if f == "=TABLE" or (ws.title == "Sensitivity" and not isinstance(c.value, (str, int, float))):
                    f = table_f
                val = values.get((ws.title.upper(), c.coordinate), c.value) if f else c.value
                if isinstance(val, float) and val != val:  # NaN
                    val = None
                color = (c.font.color.rgb if c.font and c.font.color and isinstance(c.font.color.rgb, str) else "")[-6:]
                role = "input" if color == BLUE else "link" if color == GREEN else "formula" if f else "label"
                cells.append({"c": c.column, "ref": c.coordinate, "v": val if not isinstance(val, (date,)) else str(val), "f": f,
                              "fmt": c.number_format, "role": role, "bold": bool(c.font and c.font.bold),
                              "fill": (c.fill.fgColor.rgb[-6:] if c.fill and c.fill.fill_type == "solid" else None)})
            if cells:
                rows.append({"r": row[0].row, "cells": cells})
        sheets.append({"name": ws.title, "max_col": ws.max_column, "rows": rows,
                       "widths": {k: v.width for k, v in ws.column_dimensions.items() if v.width}})
    out = {"scenario": scenario, "meta": meta, "sheets": sheets,
           "status": values.get(tuple(x.upper() if i == 0 else x for i, x in enumerate(meta["status_cell"].split("!"))))}
    for k in [k for k in _GRID_CACHE if k[1] != stamp]:  # a new day's data replaces yesterday's three scenarios
        del _GRID_CACHE[k]
    _GRID_CACHE[key] = out
    return out


def _headcount(wh, months):
    out = []
    for m in months:
        end = _months_after(m, 1) + "-01"
        out.append(int(wh.query("""select count(*) as n from staging.stg_payroll__employees
                                   where hire_date < ?::date and (termination_date is null or termination_date < hire_date
                                                                  or termination_date >= ?::date)""", [end, end])[0]["n"]))
    return out


def _guide(ws, actual, fc):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120
    lines = [
        (f"{COMPANY} · three-statement operating model", _font(bold=True, size=14)),
        (f"Actuals {actual[1]} to {actual[-1]} from the ledger; forecast {fc[0]} to {fc[-1]} by formula.", _font(MUTED)),
        ("", None),
        ("How it links", _font(bold=True)),
        ("1. Assumptions: pick Base, Upside or Downside in C4. The Live column picks that scenario's drivers with INDEX/MATCH.", None),
        ("2. P&L: revenue grows from last month by the growth driver; costs are driver percentages or roll forward with growth and the hiring plan.", None),
        ("3. Balance sheet: receivables and payables from DSO and DPO; fixed assets roll forward with capex and depreciation; "
         "retained earnings roll forward with P&L net income; cash comes from the cash flow statement.", None),
        ("4. Cash flow (indirect): net income + depreciation - change in receivables + change in payables and payroll liabilities, "
         "less capex, plus any equity raised. Closing cash feeds the balance sheet.", None),
        ("5. KPIs: burn, runway, ARR, CAC payback and margins, all formulas on the statements.", None),
        ("6. Checks: the balance sheet balances, cash ties, net income ties to the ledger, retained earnings roll forward: TRUE in every month.", None),
        ("", None),
        ("Colour code", _font(bold=True)),
        ("Blue: input (a ledger actual or an assumption you can change).", _font(BLUE)),
        ("Black: formula on the same sheet.", _font(BLACK)),
        ("Green: formula linking to another sheet.", _font(GREEN)),
        ("Shaded columns are forecast months; they contain no typed numbers.", None),
        ("", None),
        ("Simplifications", _font(bold=True)),
        ("One cash account and one company (entities consolidated); no debt drawn; no interest income; tax shown at the cash rate "
         "(US losses carried forward shelter profits). Edit the drivers or enter an equity raise on the Cash flow sheet to test a plan.", None),
    ]
    for k, (text, font) in enumerate(lines, 1):
        c = ws.cell(k, 1, text)
        if font:
            c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--out", default=str(ROOT / "reports" / "echo_labs_operating_model.xlsx"))
    a = p.parse_args()
    data, meta = build()
    from pathlib import Path
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_bytes(data)
    print(f"wrote {a.out}: {meta}")


if __name__ == "__main__":
    main()
