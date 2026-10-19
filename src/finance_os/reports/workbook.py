"""The finance workbook (.xlsx): every statement and schedule a finance team keeps in Excel, built from the warehouse.

    python -m finance_os.reports.workbook --period quarter      # → reports/finance_workbook.xlsx

Sheets: Income statement, Balance sheet, Cash flow (from get_financial_statements, one column per period), KPIs,
Budget vs actual (last closed month), AR aging, the 18-month forecast, unit economics, cost structure, price/volume/mix,
runway sensitivity and the five-year plan. Numbers are values, formatted the way
an analyst would: thousands separators, negatives in red brackets, percentages, frozen headers, bold totals.
"""
import argparse
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finance_os.app import tools
from finance_os.paths import ROOT

USD = '#,##0;[Red](#,##0);"–"'
PCT = '0.0%;[Red]-0.0%;"–"'
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
SECTION_FILL = PatternFill("solid", fgColor="F3F4F6")
TOP = Border(top=Side(style="thin", color="9CA3AF"))
DOUBLE = Border(top=Side(style="thin", color="9CA3AF"), bottom=Side(style="double", color="9CA3AF"))


def _header(ws, labels, title):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(labels)
    for c in ws[2]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(horizontal="left" if c.column == 1 else "right")
    ws.freeze_panes = "B3"


def _widths(ws, first=34, rest=13):
    ws.column_dimensions["A"].width = first
    for i in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = rest


def _statement(ws, sheet, rows, periods, subtitle):
    _header(ws, ["", *periods], f"Echo Labs · {sheet} · {subtitle}")
    for r in (r for r in rows if r["sheet"] == sheet):
        ws.append([("    " if r["style"] == "item" else "") + r["line"], *[r[p] for p in periods]])
        row = ws[ws.max_row]
        for c in row[1:]:
            c.number_format = PCT if r["unit"] == "ratio" else USD
        if r["style"] == "header":
            for c in row:
                c.fill, c.font = SECTION_FILL, Font(bold=True)
        elif r["style"] in ("subtotal", "total"):
            for c in row:
                c.font, c.border = Font(bold=True), DOUBLE if r["style"] == "total" else TOP
        elif r["style"] in ("ratio", "check"):
            for c in row:
                c.font = Font(italic=True, color="6B7280")
    _widths(ws)


def _table(ws, title, rows, columns):
    """columns: [(key, label, format or None)]"""
    _header(ws, [l for _, l, _ in columns], title)
    for r in rows:
        ws.append([r.get(k) for k, _, _ in columns])
        for c, (_, _, fmt) in zip(ws[ws.max_row], columns):
            if fmt:
                c.number_format = fmt
    _widths(ws, 22, 15)


def build(wh, period="month", periods=12):
    fs = tools.get_financial_statements(wh, period=period, periods=periods)
    end = fs["end_month"][:7]
    sub = f"{period}ly through {end}" if period != "month" else f"monthly through {end}"
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in ("Income statement", "Balance sheet", "Cash flow"):
        _statement(wb.create_sheet(sheet), sheet, fs["rows"], fs["periods"], sub)

    kpi = wh.query("""select strftime(month, '%Y-%m') as month, metric_name, value from marts.fct_kpi_monthly
                      where month <= ?::date order by month""", [fs["end_month"]])
    months = sorted({r["month"] for r in kpi})[-12:]
    by = {(r["metric_name"], r["month"]): r["value"] for r in kpi}
    ws = wb.create_sheet("KPIs")
    _header(ws, ["", *months], f"Echo Labs · KPIs · monthly through {end}")
    ratio = {"gross_margin_pct", "nrr", "grr", "logo_churn_rate", "gross_mrr_churn_rate", "cei"}
    plain = {"runway_months", "headcount", "active_customers", "new_customers", "churned_customers", "dso_days"}
    for name in dict.fromkeys(r["metric_name"] for r in kpi):
        ws.append([name.replace("_", " ").capitalize(), *[by.get((name, m)) for m in months]])
        for c in ws[ws.max_row][1:]:
            c.number_format = PCT if name in ratio else ("0.0" if name in plain else USD)
    _widths(ws)

    bva = tools.get_budget_vs_actual(wh)
    _table(wb.create_sheet("Budget vs actual"), f"Budget vs actual · {end}", bva["rows"],
           [("department_id", "Department", None), ("account_name", "Account", None), ("actual", "Actual", USD),
            ("budget", "Budget", USD), ("variance", "Variance", USD), ("variance_pct", "Variance %", PCT),
            ("ytd_actual", "YTD actual", USD), ("ytd_budget", "YTD budget", USD), ("fy_forecast", "FY forecast", USD)])

    aging = wh.query("""select invoice_ref, customer_id, segment, invoice_type, balance, days_past_due, aging_bucket
                        from marts.fct_ar_aging_snapshot
                        where snapshot_date = (select max(snapshot_date) from marts.fct_ar_aging_snapshot)
                        order by days_past_due desc""")
    _table(wb.create_sheet("AR aging"), "Open receivables today", aging,
           [("invoice_ref", "Invoice", None), ("customer_id", "Customer", None), ("segment", "Segment", None),
            ("invoice_type", "Type", None), ("balance", "Balance", USD), ("days_past_due", "Days past due", "0"),
            ("aging_bucket", "Bucket", None)])

    fc = wh.query("""select scenario, strftime(month, '%Y-%m') as month, revenue, cogs, opex, ebitda, closing_cash,
                            closing_cash_with_raise
                     from marts.fct_forecast_monthly where is_latest order by scenario, month""")
    _table(wb.create_sheet("Forecast"), "18-month forecast by scenario", fc,
           [("scenario", "Scenario", None), ("month", "Month", None), ("revenue", "Revenue", USD), ("cogs", "COGS", USD),
            ("opex", "Opex", USD), ("ebitda", "EBITDA", USD), ("closing_cash", "Cash, no new money", USD),
            ("closing_cash_with_raise", "Cash with planned raise", USD)])
    saas = tools.get_saas_metrics(wh, 12)["rows"]
    _table(wb.create_sheet("Unit economics"), "SaaS unit economics (CAC trailing 3 months)", saas,
           [("month", "Month", None), ("arr", "ARR", USD), ("new_customers", "New customers", "0"), ("sm_spend", "S&M spend", USD),
            ("cac", "CAC", USD), ("arpa", "ARPA", USD), ("ltv", "LTV", USD), ("ltv_to_cac", "LTV:CAC", "0.0x"),
            ("cac_payback_months", "CAC payback (months)", "0.0"), ("revenue_growth_yoy", "Growth YoY", PCT),
            ("ebitda_margin", "EBITDA margin", PCT), ("rule_of_40", "Rule of 40", PCT), ("free_cash_flow", "Free cash flow", USD)])
    cost = tools.get_cost_behavior(wh, 12)["rows"]
    _table(wb.create_sheet("Cost structure"), "Fixed vs variable costs and breakeven", cost,
           [("month", "Month", None), ("revenue", "Revenue", USD), ("variable_costs", "Variable costs", USD),
            ("fixed_costs", "Fixed costs", USD), ("contribution_margin", "Contribution margin", USD), ("cm_ratio", "CM %", PCT),
            ("breakeven_revenue", "Breakeven revenue", USD), ("margin_of_safety", "Margin of safety", PCT),
            ("operating_leverage", "Operating leverage", "0.00")])
    pvm = tools.get_revenue_pvm(wh)
    _table(wb.create_sheet("Price volume mix"), f"Usage revenue bridge · {str(pvm['month'])[:7]} vs prior month", pvm["rows"],
           [("product_id", "Product", None), ("revenue_prev", "Prior revenue", USD), ("revenue", "Revenue", USD),
            ("price_effect", "Price", USD), ("volume_effect", "Volume", USD), ("mix_effect", "Mix", USD), ("revenue_change", "Change", USD)])
    sens = tools.get_runway_sensitivity(wh)
    cols = list(sens["rows"][0])
    _table(wb.create_sheet("Runway sensitivity"), "Months of runway: revenue growth (rows) × opex change (columns)", sens["rows"],
           [(c, "Revenue growth / month" if c == "revenue_growth_monthly" else c, PCT if c == "revenue_growth_monthly" else "0") for c in cols])
    lrp = tools.get_long_range_plan(wh)["rows"]
    _table(wb.create_sheet("Five-year plan"), "Long-range plan", lrp,
           [("year", "Year", "0"), ("revenue", "Revenue", USD), ("revenue_growth", "Growth", PCT), ("gross_margin_pct", "GM %", PCT),
            ("sales_marketing", "S&M", USD), ("research_development", "R&D", USD), ("general_admin", "G&A", USD),
            ("ebitda", "EBITDA", USD), ("ebitda_margin", "EBITDA %", PCT), ("capex", "Capex", USD),
            ("free_cash_flow", "Free cash flow", USD), ("financing", "Financing", USD), ("closing_cash", "Closing cash", USD)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--period", choices=["month", "quarter", "year"], default="month")
    p.add_argument("--periods", type=int, default=12)
    a = p.parse_args()
    out = ROOT / "reports" / "finance_workbook.xlsx"
    out.parent.mkdir(exist_ok=True)
    out.write_bytes(build(tools.Warehouse(), a.period, a.periods))
    print(out)


if __name__ == "__main__":
    main()
