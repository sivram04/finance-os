"""Recurring reports: cadence rules, and every report builds from the warehouse."""
from datetime import date

import pytest

from finance_os.reports import scheduled


def test_cadences_decide_the_period_each_report_covers():
    closed = date(2026, 9, 1)
    assert scheduled.period_key("daily", date(2026, 10, 18), closed) == "2026-10-18"
    assert scheduled.period_key("weekly:mon", date(2026, 10, 19), closed) == "2026-W43"     # a Monday
    assert scheduled.period_key("weekly:mon", date(2026, 10, 18), closed) is None            # a Sunday
    assert scheduled.period_key("monthly:on_close", date(2026, 10, 18), closed) == "2026-09"  # once per closed month


@pytest.mark.skipif(not scheduled.WAREHOUSE.exists(), reason="build the warehouse first")
def test_every_report_builds():
    import duckdb
    con = duckdb.connect(str(scheduled.WAREHOUSE), read_only=True)
    try:
        for name, build in scheduled.REPORTS.items():
            text = build(con, date(2026, 10, 19))
            assert text.startswith("#") and "None" not in text.splitlines()[0], name
    finally:
        con.close()


def test_three_statements_tie_out():
    """Balance sheet balances, and the cash flow ends at the balance sheet's cash, every period."""
    from finance_os.app.tools import Warehouse, get_financial_statements
    for period in ("month", "quarter", "year"):
        fs = get_financial_statements(Warehouse(), period=period, periods=6)
        line = {(r["sheet"], r["line"]): r for r in fs["rows"]}
        for p in fs["periods"]:
            assert abs(line[("Balance sheet", "Check (should be 0)")][p]) < 0.05, (period, p)
            assert abs(line[("Cash flow", "Closing cash")][p] - line[("Balance sheet", "Cash - operating")][p]) < 0.05


def test_finance_workbook_has_every_sheet():
    import io
    from openpyxl import load_workbook
    from finance_os.app.tools import Warehouse
    from finance_os.reports import workbook
    wb = load_workbook(io.BytesIO(workbook.build(Warehouse(), "quarter", 4)))
    assert {"Income statement", "Balance sheet", "Cash flow", "KPIs", "Unit economics", "Cost structure",
            "Price volume mix", "Runway sensitivity", "Five-year plan"} <= set(wb.sheetnames)
